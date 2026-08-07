"""
Generación de Excel (.xlsx) a partir de la data de los colaboradores:
  - inventario_punto_bytes(): el inventario de UN punto, para que el colaborador lo reenvíe (bot).
  - maestro_bytes(): export maestro para instituciones; completa=True añade el censo de personas
    (DATOS PERSONALES → solo admin con login).
Usa openpyxl (ya instalado). Sin estado: lee de db en cada llamada.
"""
import io

import openpyxl
from openpyxl.styles import Font, PatternFill

import db

_CLASE_LABEL = {"oferta": "Tiene (para dar)", "stock": "Tiene (stock privado)", "demanda": "Necesita"}


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _encabezado(ws, headers):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B2A4A")
    ws.freeze_panes = "A2"


def _ajustar(ws):
    for col in ws.columns:
        ancho = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(ancho + 2, 10), 50)


def inventario_punto_bytes(punto_id, nombre_punto=""):
    """Excel del inventario de UN punto (sin datos personales)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario"
    _encabezado(ws, ["Situación", "Categoría", "Insumo", "Cantidad", "Descripción", "Cargado"])
    for r in db.inventario_de_punto(punto_id):
        clase = r.get("clase") or ("oferta" if r.get("tipo_punto") in db.TIPOS_OFERTA else "demanda")
        ws.append([_CLASE_LABEL.get(clase, clase), r.get("categoria"), r.get("item"),
                   r.get("cantidad"), r.get("detalle"), (r.get("creado_en") or "")[:16]])
    _ajustar(ws)
    return _bytes(wb)


_PUNTOS_COLS = [
    ("punto", "Punto"), ("tipo_punto", "Tipo"), ("rol", "Rol"), ("responsable", "Responsable"),
    ("telefono", "Contacto"), ("ubicacion", "Ubicación"), ("lat", "Lat"), ("lon", "Lon"),
    ("horario", "Horario"), ("tipo_inmueble", "Inmueble"), ("aforo", "Aforo"), ("camas", "Camas"),
    ("banos", "Baños"), ("duchas", "Duchas"), ("agua_potable", "Agua"), ("electricidad", "Luz"),
    ("planta", "Planta"), ("cocina", "Cocina"), ("refrigeracion", "Refrigeración"),
    ("almacenamiento", "Almacenamiento"), ("acceso_vehiculos", "Acceso vehíc."),
    ("accesibilidad", "Accesible"), ("rif", "RIF"), ("cobertura", "Cobertura"),
    ("creado_en", "Registrado"),
]
_INV_COLS = [("punto", "Punto"), ("tipo_punto", "Tipo"), ("situacion", "Situación"),
             ("categoria", "Categoría"), ("item", "Insumo"), ("cantidad", "Cantidad"),
             ("detalle", "Descripción"), ("creado_en", "Cargado")]
_SEG_COLS = [("punto", "Punto"), ("tipo_punto", "Tipo"), ("item", "Insumo"), ("estado", "Estado"),
             ("match_punto", "Acopio que lo cubre"), ("actualizado_en", "Actualizado")]
_CENSO_COLS = [("punto", "Punto/refugio"), ("nombre", "Nombre"), ("cedula", "Cédula"),
               ("edad", "Edad"), ("detalle", "Detalle"), ("creado_en", "Registrado")]


def _hoja(wb, titulo, cols, filas, primera=False):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo[:31]
    _encabezado(ws, [h for _, h in cols])
    for r in filas:
        ws.append([r.get(k) for k, _ in cols])
    _ajustar(ws)


def maestro_bytes(completa=False):
    """Excel maestro para instituciones. completa=True añade 'Censo personas' (DATOS PERSONALES)."""
    wb = openpyxl.Workbook()
    _hoja(wb, "Puntos", _PUNTOS_COLS, db.export_puntos(), primera=True)
    _hoja(wb, "Inventario", _INV_COLS, db.export_inventario())
    _hoja(wb, "Seguimiento", _SEG_COLS, db.export_seguimiento())
    if completa:
        _hoja(wb, "Censo personas", _CENSO_COLS, db.export_censo())
    return _bytes(wb)
